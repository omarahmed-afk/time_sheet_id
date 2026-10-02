import argparse
import os
import sys
import datetime as dt
from zoneinfo import ZoneInfo
from pathlib import Path


import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


# ============================== CONFIG ==============================
TOKEN = os.environ.get("DEPUTY_TOKEN", "").strip()


INSTALL_URL = os.environ.get("DEPUTY_INSTALL_URL", "https://ptofthecity.na.deputy.com").rstrip("/")
API_BASE    = f"{INSTALL_URL}/api/v1"
                                   


# ---- Rules ----
NY               = ZoneInfo("America/New_York")
OT_THRESHOLD_MIN = 11                                  # keep employees 11+ minutes past shift end
TARGET_DATE      = "yesterday"                         # New York today/yesterday, a YYYY-MM-DD date, or None for previous days
OT_MODE          = "end"
CLINICS_FILTER   = []
OUTPUT_DIR       = Path(__file__).resolve().parent / "reports"
# ===================================================================


HEADERS = {
    "Authorization": f"Bearer {TOKEN}",
    "Content-Type": "application/json",
    "Accept": "application/json",
}




def query_resource(resource: str, payload: dict) -> list:
    """Call /api/v1/resource/<Name>/QUERY with pagination (500 max per request)."""
    if not TOKEN:
        raise RuntimeError("Set the DEPUTY_TOKEN environment variable before running.")
    out, start = [], 0
    while True:
        body = dict(payload, max=500, start=start)
        url = f"{API_BASE}/resource/{resource}/QUERY"
        resp = requests.post(url, json=body, headers=HEADERS, timeout=60)
        if resp.status_code == 401:
            sys.exit("Token rejected (401). It may have expired - paste a fresh one.")
        resp.raise_for_status()
        chunk = resp.json()
        if not isinstance(chunk, list):
            raise RuntimeError(f"Unexpected response from {resource}: {chunk}")
        out.extend(chunk)
        if len(chunk) < 500:
            break
        start += 500
    return out




def validate_token():
    r = requests.get(f"{API_BASE}/me", headers=HEADERS, timeout=30)
    if r.status_code == 401:
        sys.exit("Token rejected (401). The 24h window may be over - paste a fresh one.")
    r.raise_for_status()
    me = r.json()
    print(f"Token OK. Logged in as: {me.get('Name', me.get('DisplayName', '?'))}")




def target_range(target):
    """Today/yesterday in New York, a fixed date, or previous days when target is None."""
    today = dt.datetime.now(NY).date()
    if target == "today":
        start_date = end_date = today
    elif target == "yesterday":
        start_date = end_date = today - dt.timedelta(days=1)
    elif isinstance(target, str):
        start_date = end_date = dt.date.fromisoformat(target)
    elif today.weekday() == 0:
        start_date = today - dt.timedelta(days=3)
        end_date = today - dt.timedelta(days=1)
    else:
        start_date = end_date = today - dt.timedelta(days=1)
    start = dt.datetime.combine(start_date, dt.time.min, NY)
    end = dt.datetime.combine(end_date, dt.time.max, NY)
    return start_date, end_date, int(start.timestamp()), int(end.timestamp())




def to_ny(unix_ts: int) -> dt.datetime:
    return dt.datetime.fromtimestamp(unix_ts, NY)




# Only these roles appear in the report. Anything else -> row is dropped.
ROLE_MAP = {
    "physical therapist": "PT",
    "physical threapist": "PT",                 # misspelling found in the data
    "physical therapy assistant": "PTA",
    "patients care coordinator": "PCC",
    "patients care coordinator (pcc)": "PCC",
    "aide": "PT Aide",
}




def map_role(name: str):
    """Return the short code for an approved role, or None to drop the row."""
    return ROLE_MAP.get((name or "").strip().lower())




def fetch_timesheets(start_unix, end_unix):
    """Fetch the day's timesheets with their scheduled Roster attached."""
    search = {
        "f1": {"field": "StartTime", "data": start_unix, "type": "ge"},
        "f2": {"field": "StartTime", "data": end_unix, "type": "le"},
    }
    try:
        timesheets = query_resource("Timesheet", {"search": search, "join": ["RosterObject"]})
    except RuntimeError:
        timesheets = query_resource("Timesheet", {"search": search})
    print(f"  Found {len(timesheets)} timesheets.")


    roster_ids = sorted({t["Roster"] for t in timesheets
                         if t.get("Roster") and not t.get("RosterObject")})
    rosters = {}
    for i in range(0, len(roster_ids), 200):
        rp = {"search": {"r1": {"field": "Id", "data": roster_ids[i:i + 200], "type": "in"}}}
        for r in query_resource("Roster", rp):
            rosters[r["Id"]] = r
    for t in timesheets:
        t["_roster"] = t.get("RosterObject") or rosters.get(t.get("Roster"))
    print(f"  Found {len(rosters)} scheduled rosters (separate query).")
    return timesheets




def evaluate_timesheet(t):
    """Return decision, explanation and overtime seconds; fail closed on incomplete data."""
    meta = t.get("_DPMetaData") or {}
    ou = meta.get("OperationalUnitInfo") or {}
    facility = ou.get("CompanyName", "")
    if (not facility or "pediatric" in facility.lower()
            or map_role(ou.get("OperationalUnitName", "")) is None
            or (CLINICS_FILTER and facility not in CLINICS_FILTER)):
        return "SKIP", "Outside configured clinics/roles or missing metadata", None
    roster = t.get("_roster") or {}
    start, end = t.get("StartTime"), t.get("EndTime")
    rs, re = roster.get("StartTime"), roster.get("EndTime")
    if (not all(isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0
                for v in (start, end, rs, re)) or end <= start or re <= rs):
        return "REVIEW", "Missing or invalid actual/scheduled times", None
    overtime = end - re if OT_MODE == "end" else (end - start) - (re - rs)
    if overtime >= OT_THRESHOLD_MIN * 60:
        return "OVERTIME", "Overtime threshold reached; no approval", overtime
    if not t.get("Id"):
        return "REVIEW", "Missing timesheet ID", overtime
    if any(t.get(key) for key in ("Discarded", "isInProgress", "IsLeave", "Disputed",
                                  "Exported", "PayStaged")):
        return "REVIEW", "Leave, in-progress, discarded, disputed or locked record", overtime
    if t.get("ValidationFlag") != 0:
        return "REVIEW", "Validation warning or missing validation status (check absence)", overtime
    if end > dt.datetime.now(NY).timestamp():
        return "REVIEW", "Future end time", overtime
    if t.get("TimeApproved") or t.get("PayRuleApproved"):
        return "SKIP", "Already time/pay approved", overtime
    if t.get("TimeApproved") is not False or t.get("PayRuleApproved") is not False:
        return "REVIEW", "Unknown approval state", overtime
    return "READY", "Below overtime threshold", overtime


def overtime_row(t, overtime):
    meta = t.get("_DPMetaData") or {}
    ou = meta.get("OperationalUnitInfo") or {}
    roster = t["_roster"]
    return [
        (meta.get("EmployeeInfo") or {}).get("DisplayName", ""),
        ou.get("CompanyName", ""), map_role(ou.get("OperationalUnitName", "")),
        to_ny(t["StartTime"]).strftime("%m/%d/%Y"), int(round(overtime / 60)),
        to_ny(roster["EndTime"]).strftime("%I:%M %p"),
        t.get("SupervisorComment") or t.get("EmployeeComment") or roster.get("Comment") or "",
    ]


def make_plan(timesheets):
    rows, plan = [], []
    for t in timesheets:
        decision, reason, overtime = evaluate_timesheet(t)
        name = ((t.get("_DPMetaData") or {}).get("EmployeeInfo") or {}).get("DisplayName", "")
        plan.append([t.get("Id"), name, decision, reason])
        if decision == "OVERTIME":
            rows.append(overtime_row(t, overtime))
    rows.sort(key=lambda r: (r[1], r[0]))
    return rows, plan


def build_rows(start_unix, end_unix):
    return make_plan(fetch_timesheets(start_unix, end_unix))[0]


def approve_timesheet(timesheet_id, start_unix, end_unix):
    # Re-read immediately before approval: the timesheet/roster may have changed.
    records = query_resource("Timesheet", {"search": {
        "id": {"field": "Id", "data": timesheet_id, "type": "eq"}}})
    if len(records) != 1 or records[0].get("Id") != timesheet_id:
        return "REVIEW", "Timesheet could not be reloaded"
    t = records[0]
    rosters = query_resource("Roster", {"search": {
        "id": {"field": "Id", "data": t["Roster"], "type": "eq"}}}) if t.get("Roster") else []
    t["_roster"] = rosters[0] if len(rosters) == 1 else None
    decision, reason, _ = evaluate_timesheet(t)
    if decision != "READY":
        return decision, "Recheck: " + reason
    if not start_unix <= t["StartTime"] <= end_unix:
        return "REVIEW", "Timesheet moved outside selected dates"
    # Documented Deputy approval action; never change hours or call it for overtime.
    # https://developer.deputy.com/docs/timesheet-management-calls
    response = requests.post(f"{API_BASE}/supervise/timesheet/approve",
                             json={"intTimesheetId": timesheet_id},
                             headers=HEADERS, timeout=60)
    response.raise_for_status()
    result = response.json()
    if not isinstance(result, dict) or result.get("Id") != timesheet_id or result.get("TimeApproved") is not True:
        raise RuntimeError("Approval outcome unconfirmed; check Deputy before retrying")
    if result.get("PayRuleApproved") is True:
        return "APPROVED", "Deputy confirmed time and pay approval"
    return "TIME_APPROVED", "Deputy confirmed time approval; pay still pending"


# --------------------------- Local Excel -----------------------------
def pending_rows(timesheets, plan):
    """Include only REVIEW records still awaiting time or pay approval."""
    decisions = {record[0]: record for record in (plan or [])}
    pending = []
    def formatted(timestamp, pattern):
        return to_ny(timestamp).strftime(pattern) if timestamp else ""
    for t in timesheets:
        record = decisions.get(t.get("Id"), [t.get("Id"), "", "REVIEW", "Not evaluated"])
        if t.get("Discarded") or record[2] != "REVIEW":
            continue
        if t.get("TimeApproved") is True and t.get("PayRuleApproved") is True:
            continue
        meta = t.get("_DPMetaData") or {}
        ou = meta.get("OperationalUnitInfo") or {}
        roster = t.get("_roster") or {}
        if record[2] == "CHECK_DEPUTY":
            awaiting = "Unknown - verify in Deputy"
        elif record[2] == "TIME_APPROVED" or t.get("TimeApproved") is True:
            awaiting = "Pay approval"
        elif t.get("TimeApproved") is False:
            awaiting = "Time approval" if t.get("PayRuleApproved") is True else "Time and pay approval"
        else:
            awaiting = "Unknown - verify in Deputy"
        pending.append([
            t.get("Id"), (meta.get("EmployeeInfo") or {}).get("DisplayName", ""),
            ou.get("CompanyName", ""), ou.get("OperationalUnitName", ""),
            formatted(t.get("StartTime"), "%m/%d/%Y"),
            formatted(t.get("StartTime"), "%I:%M %p"),
            formatted(t.get("EndTime"), "%I:%M %p"),
            formatted(roster.get("EndTime"), "%I:%M %p"),
            awaiting, record[2], record[3],
        ])
    return sorted(pending, key=lambda r: (r[2], r[1]))


def write_to_sheet(rows, start_date, end_date, plan=None, output_path=None, timesheets=None):
    """Save each report as a separate local Excel workbook."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    run_stamp = dt.datetime.now(NY).strftime("%Y%m%d_%H%M%S_%f")
    output_path = output_path or OUTPUT_DIR / f"overtime_{start_date}_{end_date}_{run_stamp}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Overtime"
    ws.append([
        "EMPLOYEE NAME", "FACILITY", "ROLE", "DATE",
        "OVERTIME (MIN)", "SCHD. SHIFT END", "REASON",
    ])
    for row in rows:
        ws.append(row)
        # Names and comments are literal text, including any leading '='.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        ws.cell(ws.max_row, 5).number_format = "0"
        ws.cell(ws.max_row, 7).alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24476A")
    for column, width in zip("ABCDEFG", [28, 32, 14, 15, 20, 22, 60]):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if plan is not None:
        audit = wb.create_sheet("Approvals")
        audit.append(["TIMESHEET ID", "EMPLOYEE NAME", "STATUS", "DETAILS"])
        for record in plan:
            audit.append(record)
            for cell in audit[audit.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for column, width in zip("ABCD", [20, 30, 24, 85]):
            audit.column_dimensions[column].width = width
        audit.freeze_panes = "A2"
        audit.auto_filter.ref = audit.dimensions
    pending = wb.create_sheet("Pending")
    pending.append([
        "TIMESHEET ID", "EMPLOYEE NAME", "FACILITY", "ROLE", "DATE",
        "ACTUAL START", "ACTUAL END", "SCHD. SHIFT END", "AWAITING", "STATUS", "REASON",
    ])
    for record in pending_rows(timesheets or [], plan):
        pending.append(record)
        for cell in pending[pending.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        pending.cell(pending.max_row, 11).alignment = Alignment(wrap_text=True, vertical="top")
    for cell in pending[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24476A")
    for column, width in zip("ABCDEFGHIJK", [18, 28, 32, 30, 15, 18, 18, 22, 32, 22, 70]):
        pending.column_dimensions[column].width = width
    pending.freeze_panes = "A2"
    pending.auto_filter.ref = pending.dimensions
    try:
        wb.save(output_path)
    except PermissionError:
        sys.exit(f"Cannot save Excel report. Check folder permissions: {output_path}")
    finally:
        wb.close()
    print(f"  Saved {len(rows)} rows to: {output_path}")
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Export overtime and approve eligible Deputy timesheets.")
    parser.add_argument("--approve", action="store_true", help="Apply approvals in Deputy; default is preview only")
    parser.add_argument("--date", help="Target date: today, yesterday or YYYY-MM-DD (default: TARGET_DATE, yesterday in New York)")
    args = parser.parse_args()
    if OT_MODE not in ("end", "total") or OT_THRESHOLD_MIN <= 0:
        sys.exit("Invalid overtime mode or threshold.")
    if not TOKEN:
        sys.exit("Set the DEPUTY_TOKEN environment variable before running.")


    validate_token()
    start_date, end_date, start_unix, end_unix = target_range(args.date or TARGET_DATE)
    if start_date == end_date:
        print(f"Date: {start_date.strftime('%m/%d/%Y')} (NY time)")
    else:
        print(f"Dates: {start_date.strftime('%m/%d/%Y')} - "
              f"{end_date.strftime('%m/%d/%Y')} (NY time)")


    timesheets = fetch_timesheets(start_unix, end_unix)
    rows, plan = make_plan(timesheets)
    output_path = write_to_sheet(rows, start_date, end_date, plan, timesheets=timesheets)
    ready = sum(record[2] == "READY" for record in plan)
    print(f"  {len(rows)} overtime records; {ready} eligible for approval.")
    if not args.approve:
        print("  PREVIEW ONLY: no approvals sent. Use --approve to apply eligible approvals.")
    else:
        for record in plan:
            if record[2] != "READY":
                continue
            try:
                record[2], record[3] = approve_timesheet(record[0], start_unix, end_unix)
            except (requests.RequestException, ValueError, RuntimeError) as exc:
                record[2] = "CHECK_DEPUTY"
                record[3] = f"Stopped: {type(exc).__name__}; verify current status in Deputy before retrying"
                write_to_sheet(rows, start_date, end_date, plan, output_path, timesheets)
                sys.exit(record[3])
            write_to_sheet(rows, start_date, end_date, plan, output_path, timesheets)
            print(f"  Timesheet {record[0]}: {record[2]}")
    print("Done.")




if __name__ == "__main__":
    main()
